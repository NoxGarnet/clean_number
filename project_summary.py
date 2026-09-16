#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Generate a project-level flat summary without replacing per-file ledgers."""

from __future__ import annotations

import os
import re
import shutil
from pathlib import Path

import pandas as pd

from experiment_type import detect as detect_experiment_type


SUMMARY_COLUMNS = [
    "project_id", "test_type", "age_days", "group_key", "group_id", "mix_id",
    "group_label_cn", "group_label_en", "n_specimens", "n_measurements",
    "n_excluded", "statistical_unit", "strength_mean_MPa", "strength_std_MPa",
    "modulus_mean_GPa", "modulus_std_GPa", "quality_flags", "source_table",
]


def _value(row, name, fallback=None):
    value = row.get(name, fallback)
    return fallback if pd.isna(value) else value


def _normalise_table(table_path: Path, project_id: str) -> list[dict]:
    table = pd.read_csv(table_path)
    detected_type = detect_experiment_type(str(table_path.parents[1])).kind.value
    rows = []
    for _, source in table.iterrows():
        group_key = str(_value(source, "Group", ""))
        age_match = re.match(r"d(\d+)_", group_key)
        test_type = str(_value(source, "Test_Type", detected_type))
        flags = str(_value(source, "Quality_Flags", ""))
        record = {
            "project_id": project_id,
            "test_type": test_type,
            "age_days": int(_value(source, "Age_Days", age_match.group(1) if age_match else 0)),
            "group_key": group_key,
            "group_id": str(_value(source, "Group_ID", group_key.split("_", 1)[-1])),
            "mix_id": str(_value(source, "Mix_ID", "")),
            "group_label_cn": str(_value(source, "Label_CN", _value(source, "Label", group_key))),
            "group_label_en": str(_value(source, "Label_EN", _value(source, "Label", group_key))),
            "n_specimens": int(_value(source, "n", 0)),
            "n_measurements": int(_value(source, "n_measurements", _value(source, "n", 0))),
            "n_excluded": int(_value(source, "n_excluded", 0)),
            "statistical_unit": str(_value(source, "Statistical_Unit", "specimen")),
            "strength_mean_MPa": float(_value(source, "Strength_Mean_MPa", float("nan"))),
            "strength_std_MPa": float(_value(source, "Strength_Std_MPa", float("nan"))),
            "modulus_mean_GPa": float(_value(source, "Modulus_Mean_GPa", float("nan"))),
            "modulus_std_GPa": float(_value(source, "Modulus_Std_GPa", float("nan"))),
            "quality_flags": flags,
            "source_table": str(table_path),
        }
        for column in table.columns:
            if str(column).startswith("Factor_"):
                record[f"factor_{str(column).removeprefix('Factor_')}"] = _value(source, column, "")
        rows.append(record)
    return rows


def collect_project_summary(output_root: str, project_id: str = "") -> pd.DataFrame:
    root = Path(output_root)
    project_id = project_id or root.parent.name
    rows = []
    for table_path in sorted(root.glob("*/tables/table_statistics.csv")):
        rows.extend(_normalise_table(table_path, project_id))
    result = pd.DataFrame(rows)
    if not result.empty:
        factor_columns = sorted(column for column in result.columns if column.startswith("factor_"))
        result = result.reindex(columns=SUMMARY_COLUMNS + factor_columns)
        order = pd.Categorical(result["test_type"], ["compression", "flexural"], ordered=True)
        result = result.assign(_test_order=order).sort_values(
            ["_test_order", "age_days", "group_key"], kind="stable"
        ).drop(columns="_test_order").reset_index(drop=True)
    return result


def _status(row: pd.Series) -> str:
    flags = str(row.get("quality_flags", ""))
    if "duplicate_content" in flags:
        return "待复核：存在内容重复的原始文件"
    if int(row.get("n_excluded", 0)) > 0:
        return f"已按配置排除 {int(row['n_excluded'])} 条测量"
    return "正常"


def _write_xlsx(df: pd.DataFrame, path: Path) -> None:
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    display = pd.DataFrame({
        "实验类型": df["test_type"].map({"compression": "抗压试验", "flexural": "抗折试验"}).fillna(df["test_type"]),
        "养护龄期": df["age_days"].map(lambda value: f"{int(value)}d"),
        "配比组别": df["group_label_cn"],
        "独立试件数 n": df["n_specimens"],
        "原始测量数": df["n_measurements"],
        "强度（MPa，均值 ± 标准差）": df.apply(
            lambda row: f"{row['strength_mean_MPa']:.3f} ± {row['strength_std_MPa']:.3f}", axis=1
        ),
        "模量（GPa，均值 ± 标准差）": df.apply(
            lambda row: f"{row['modulus_mean_GPa']:.3f} ± {row['modulus_std_GPa']:.3f}", axis=1
        ),
        "数据状态": df.apply(_status, axis=1),
    })
    factor_labels = {}
    insert_at = 3
    for column in sorted(c for c in df.columns if c.startswith("factor_")):
        if df[column].notna().any() and df[column].astype(str).ne("").any():
            display.insert(insert_at, factor_labels.get(column, column), df[column])
            insert_at += 1

    temporary = path.with_name(path.stem + ".tmp.xlsx")
    with pd.ExcelWriter(temporary, engine="openpyxl") as writer:
        display.to_excel(writer, sheet_name="导师汇报表", index=False)
        df.to_excel(writer, sheet_name="数值明细", index=False)
        pd.DataFrame([
            ["统计单位", "分组统计使用独立试件；抗压两个端部先在试件内平均。"],
            ["逐文件明细", "各试验输出目录中的 summary.csv 保持为逐文件处理账本。"],
            ["重复文件", "内容哈希相同的文件只标记、不自动删除；排除必须写入 processing_overrides.yaml。"],
        ], columns=["项目", "说明"]).to_excel(writer, sheet_name="数据说明", index=False)

    workbook = load_workbook(temporary)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.sheet_view.showGridLines = False
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for column in sheet.columns:
            values = [str(cell.value or "") for cell in column]
            width = min(max(max(map(len, values)) + 2, 10), 48)
            sheet.column_dimensions[column[0].column_letter].width = width
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    workbook.save(temporary)
    os.replace(temporary, path)


def generate_project_summary(
    output_root: str,
    project_id: str = "",
    summary_basename: str = "",
) -> tuple[str, str] | None:
    """Write project_summary.csv/xlsx beside experiment output directories."""
    root = Path(output_root)
    root.mkdir(parents=True, exist_ok=True)
    df = collect_project_summary(str(root), project_id=project_id)
    if df.empty:
        return None

    csv_path = root / "project_summary.csv"
    csv_temp = root / "project_summary.tmp.csv"
    df.to_csv(csv_temp, index=False, encoding="utf-8-sig")
    os.replace(csv_temp, csv_path)

    xlsx_path = root / "project_summary.xlsx"
    _write_xlsx(df, xlsx_path)

    if summary_basename:
        alias_csv = root / f"{summary_basename}.csv"
        alias_xlsx = root / f"{summary_basename}.xlsx"
        alias_csv_temp = root / f"{summary_basename}.tmp.csv"
        df.to_csv(alias_csv_temp, index=False, encoding="utf-8-sig")
        os.replace(alias_csv_temp, alias_csv)
        shutil.copy2(xlsx_path, alias_xlsx)
    return str(csv_path), str(xlsx_path)
